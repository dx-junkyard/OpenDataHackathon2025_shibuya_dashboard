"""Data 整形支援 Agent (F-09): proposes a shaping plan; execution is deterministic.

The model only fills a finite plan (which table, which header row, which
columns, which filter, which aggregation). Everything else — extracting the
tables from CSV / Excel / PDF / HTML, unpivoting a wide table, parsing 和暦
periods — is done by `table_rows` and `apply_plan` without the model, so the
user can check the steps and before/after samples before adopting (F-21).
"""
from __future__ import annotations

import csv
import io
import os
import re
import unicodedata
from collections import OrderedDict
from html.parser import HTMLParser
from typing import Any, Dict, List, Optional

from .. import jpdate, llm
from ..observe import observed

AGGREGATIONS = ["count_rows", "sum", "last_value"]
PERIODS = ["month", "year", "as_is"]
FILE_FORMATS = ("csv", "xlsx", "pdf", "html")

# Columns the system adds; the plan may refer to them like any other column.
UNPIVOT_PERIOD = "期間"
UNPIVOT_VALUE = "値"
SNAPSHOT_COLUMN = "取得時点"

SYSTEM = f"""あなたはオープンデータを KPI の時系列データへ整形する手順を提案します。
与えられるのは KPI と、ファイルから取り出した表（番号つき、各表の先頭行）です。次の有限の手順だけで表せる整形案を作ります。
- table_index: 使う表の番号。
- header_row: その表で列名が書かれている行の番号（0 始まり）。それより上の行（表題・注記）は使いません。
- unpivot_columns: 年度や月が列方向に並ぶ横長の表なら、期間を表す列名をすべて並べます。各行がそれらの列ごとの行に展開され、
  列名が「{UNPIVOT_PERIOD}」列、セルの値が「{UNPIVOT_VALUE}」列になります（date_column="{UNPIVOT_PERIOD}"、value_column="{UNPIVOT_VALUE}" とします）。縦長の表なら空配列。
- date_column: 期間を表す列（日付・年月・年度。和暦も可）。表に期間の列がなく「{SNAPSHOT_COLUMN}」列がある場合はそれを使います（取得した時点ごとの値になる）。
- value_column: 値の列。aggregation が count_rows の場合は空文字。
- filter_column / filter_value: 行の絞り込み（例: 区名・地域が「渋谷区」、項目名が KPI に当たる行）。不要なら両方空文字。
- aggregation: {AGGREGATIONS} のいずれか。count_rows=期間ごとの行数, sum=期間ごとの合計, last_value=期間ごとの最後の値。
- period: {PERIODS} のいずれか。month=YYYY-MM に丸める, year=YYYY に丸める, as_is=そのまま。
- cumulative: 期間ごとの値を累計にするなら true（例: 設置日からの累計設置数）。
- steps: 利用者向けに、何をするかを日本語で 2〜4 行。
列名は与えられたものだけを使います。整形できない場合は feasible=false にして理由を steps に書きます。"""

SCHEMA = llm.obj(
    {
        "feasible": {"type": "boolean"},
        "table_index": {"type": "integer"},
        "header_row": {"type": "integer"},
        "unpivot_columns": {"type": "array", "items": llm.STR},
        "date_column": llm.STR,
        "value_column": llm.STR,
        "filter_column": llm.STR,
        "filter_value": llm.STR,
        "aggregation": llm.enum(AGGREGATIONS),
        "period": llm.enum(PERIODS),
        "cumulative": {"type": "boolean"},
        "steps": {"type": "array", "items": llm.STR},
    }
)


class ShapingError(ValueError):
    pass


# ---------------------------------------------------------------- reading tables
# A table is {"name": str, "rows": [[cell, ...], ...]} with the raw grid,
# header rows included; `table_rows` turns one into dict rows for a plan.


def _cell(v: Any) -> str:
    if v is None:
        return ""
    return " ".join(str(v).split())


class _HtmlTables(HTMLParser):
    """<table> grids with colspan / rowspan expanded (nested tables are flattened into cells)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: List[Dict[str, Any]] = []
        self._stack: List[Dict[str, Any]] = []
        self._cell: Optional[List[str]] = None
        self._span = (1, 1)
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ("script", "style"):
            self._skip += 1
        elif tag == "table":
            self._stack.append({"rows": [], "pending": {}, "caption": [], "in_caption": False})
        elif not self._stack:
            return
        elif tag == "caption":
            self._stack[-1]["in_caption"] = True
        elif tag == "tr":
            self._stack[-1]["rows"].append([])
        elif tag in ("td", "th"):
            if not self._stack[-1]["rows"]:
                self._stack[-1]["rows"].append([])
            self._cell = []
            self._span = (_int(a.get("colspan")), _int(a.get("rowspan")))
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_data(self, data):
        if self._skip or not self._stack:
            return
        if self._cell is not None:
            self._cell.append(data)
        elif self._stack[-1]["in_caption"]:
            self._stack[-1]["caption"].append(data)

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif not self._stack:
            return
        elif tag == "caption":
            self._stack[-1]["in_caption"] = False
        elif tag in ("td", "th") and self._cell is not None:
            self._place(_cell("".join(self._cell)))
            self._cell = None
        elif tag == "table":
            t = self._stack.pop()
            if any(any(c for c in r) for r in t["rows"]):
                self.tables.append({"name": _cell("".join(t["caption"])), "rows": _rectangular(t["rows"])})

    def _place(self, text: str) -> None:
        t = self._stack[-1]
        row_i = len(t["rows"]) - 1
        row = t["rows"][row_i]
        colspan, rowspan = self._span
        for _ in range(colspan):
            while (row_i, len(row)) in t["pending"]:
                row.append(t["pending"].pop((row_i, len(row))))
            col = len(row)
            row.append(text)
            for k in range(1, rowspan):
                t["pending"][(row_i + k, col)] = text

    def close(self):
        super().close()
        for t in self._stack:  # unclosed tables
            self.tables.append({"name": "", "rows": _rectangular(t["rows"])})


def _int(v: Optional[str]) -> int:
    try:
        return max(1, min(int(v or 1), 100))
    except ValueError:
        return 1


def _rectangular(rows: List[List[str]]) -> List[List[str]]:
    rows = [r for r in rows if any(c for c in r)]
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


def _pdf_tables(path: str) -> List[Dict[str, Any]]:
    import pdfplumber

    out = []
    try:
        with pdfplumber.open(path) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                for i, t in enumerate(page.extract_tables(), start=1):
                    rows = _rectangular([[_cell(c) for c in r] for r in t])
                    if rows:
                        out.append({"name": f"{page_no} ページ目の表 {i}", "rows": rows})
    except Exception as exc:  # pdfminer raises many unrelated types for broken files
        raise ShapingError(f"PDF を読めませんでした（{type(exc).__name__}）") from exc
    if not out:
        raise ShapingError("PDF から罫線のある表が見つかりませんでした（画像の PDF や罫線のない表は未対応）")
    return out


def read_tables(path: str, file_format: str) -> List[Dict[str, Any]]:
    if file_format == "csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            return [{"name": "", "rows": [list(r) for r in csv.reader(f)]}]
    if file_format == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        out = [{"name": ws.title, "rows": _rectangular([[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)])}
               for ws in wb.worksheets]
        return [t for t in out if t["rows"]] or [{"name": "", "rows": []}]
    if file_format == "html":
        with open(path, "rb") as f:
            raw = f.read()
        text = None
        for enc in ("utf-8-sig", "cp932"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ShapingError("HTML の文字コードを判定できませんでした")
        p = _HtmlTables()
        p.feed(text)
        p.close()
        if not p.tables:
            raise ShapingError("HTML に表（table 要素）が見つかりませんでした")
        return p.tables
    if file_format == "pdf":
        return _pdf_tables(path)
    raise ShapingError(f"未対応の形式です: {file_format}")


def _unique_header(cells: List[str]) -> List[str]:
    out: List[str] = []
    for i, c in enumerate(cells):
        name = c or f"列{i + 1}"
        while name in out:
            name = f"{name}_{i + 1}"
        out.append(name)
    return out


def table_rows(tables: List[Dict[str, Any]], plan: Dict[str, Any]) -> List[Dict[str, str]]:
    """Select the plan's table and header row, unpivot if asked, and return dict rows."""
    ti = int(plan.get("table_index") or 0)
    if not 0 <= ti < len(tables):
        raise ShapingError(f"表 {ti} はありません（表は {len(tables)} 個）")
    grid = tables[ti]["rows"]
    hr = int(plan.get("header_row") or 0)
    if not 0 <= hr < len(grid):
        raise ShapingError(f"列名の行 {hr} は表の範囲外です")
    header = _unique_header([_cell(c) for c in grid[hr]])
    rows = [dict(zip(header, list(r) + [""] * (len(header) - len(r)))) for r in grid[hr + 1:]]
    cols = [c for c in plan.get("unpivot_columns") or [] if c]
    if not cols:
        return rows
    missing = [c for c in cols if c not in header]
    if missing:
        raise ShapingError(f"列が見つかりません: {'・'.join(missing)}")
    if UNPIVOT_PERIOD in header or UNPIVOT_VALUE in header:
        raise ShapingError(f"表に「{UNPIVOT_PERIOD}」「{UNPIVOT_VALUE}」という列が既にあるため展開できません")
    keep = [h for h in header if h not in cols]
    return [{**{k: r[k] for k in keep}, UNPIVOT_PERIOD: c, UNPIVOT_VALUE: r[c]} for r in rows for c in cols]


def describe_tables(tables: List[Dict[str, Any]], max_rows: int = 8) -> str:
    parts = []
    for i, t in enumerate(tables[:20]):
        buf = io.StringIO()
        csv.writer(buf).writerows(t["rows"][:max_rows])
        parts.append(f"表 {i}（{t['name'] or '名前なし'}、{len(t['rows'])} 行）:\n{buf.getvalue()}")
    if len(tables) > 20:
        parts.append(f"（ほか {len(tables) - 20} 個の表は省略）")
    return "\n".join(parts)


def profile_columns(rows: List[Dict[str, str]], max_distinct: int = 12) -> Dict[str, Any]:
    """Per-column summary over ALL rows: distinct values for low-cardinality columns."""
    out: Dict[str, Any] = {}
    for col in (rows[0].keys() if rows else []):
        values = [r.get(col, "") for r in rows]
        distinct = sorted({v for v in values if v != ""})
        out[col] = {"distinct_count": len(distinct), "values": distinct if len(distinct) <= max_distinct else [],
                    "contains_target": TARGET_AREA in distinct,
                    "sample_others": [v for v in distinct if v != TARGET_AREA][:5]}
    return out


def _prompt(kpi: Dict[str, Any], tables: List[Dict[str, Any]], note: str) -> str:
    return (f"KPI: {kpi['name']}（単位: {kpi['unit']}、測り方: {kpi['measurement']}）\n\n"
            f"{describe_tables(tables)}" + (f"\n\n{note}" if note else ""))


@observed("data-shaping-agent")
def propose_plan(kpi: Dict[str, Any], tables: List[Dict[str, Any]], note: str = "") -> Dict[str, Any]:
    """
    probe-agent:
      role: データ整形計画をAI候補として提案する入口
      capability: 根拠付きKPIダッシュボードを公開する
      system_purpose: 公開可能なKPI時系列へ向けて、元データから必要な整形手順を候補として提示する
      probe_value: データ整形もAI候補であり人が採用するという原則の適用先として、提案内容と警告の妥当性を確認しやすい
      element_type: capability
      operation_kind: analysis
      consumers: [apps.api.app.routes.data.propose_shaping]
      state_effects: [external-api]
    """
    return llm.complete_json(SYSTEM, _prompt(kpi, tables, note), SCHEMA, "shaping_plan")


TARGET_AREA = "渋谷区"


def plan_warnings(plan: Dict[str, Any], profile: Dict[str, Any], snapshot_count: Optional[int] = None) -> List[str]:
    """Deterministic checks shown next to an AI shaping plan before adoption.

    A probe-agent shadow comparison (System 6) showed the model sometimes omits
    the area filter, and that a richer prompt did not fix it reliably, so the
    guard is rule-based rather than left to the model.

    probe-agent:
      role: 整形計画のリスクや欠落条件を警告として抽出する検証関数
      capability: 根拠付きKPIダッシュボードを公開する
      system_purpose: 不完全なデータ整形や誤読を防ぎ、人が公開判断する前に注意点を明示する
      probe_value: 欠損や条件不足を可視化する支援要素であり、ダッシュボードの誤解防止ガードレールとして重要
      element_type: supporting
      operation_kind: validation
      consumers: [apps.api.app.routes.data.propose_shaping]
      state_effects: [none]
    """
    out = []
    filtered_to_target = (plan.get("filter_value") or "").strip() == TARGET_AREA and bool(plan.get("filter_column"))
    for col, p in profile.items():
        if filtered_to_target or not p.get("contains_target", TARGET_AREA in p["values"]):
            continue
        if p["distinct_count"] > 1 and plan.get("filter_column") != col:
            others = p["values"] and [v for v in p["values"] if v != TARGET_AREA] or p.get("sample_others", [])
            more = "など" if p["distinct_count"] - 1 > len(others) else ""
            out.append(f"「{col}」に{TARGET_AREA}以外（{'・'.join(others)}{more}）の行が含まれていますが、絞り込みがありません")
    date_col = plan.get("date_column") or ""
    if date_col == SNAPSHOT_COLUMN and snapshot_count is not None and snapshot_count < 2:
        out.append(f"「{SNAPSHOT_COLUMN}」が 1 回分しかないため、推移ではなく取得時点の値 1 点になります（再取得のたびに点が増えます）")
    if date_col and date_col in profile and plan.get("period") == "year":
        vals = profile[date_col]["values"] or profile[date_col].get("sample_others", [])
        if any("年度" in v for v in vals):
            out.append("年度（4月〜翌3月）を西暦の年（例: 令和5年度→2023）として表示します")
    return out


def _period(value: str, period: str) -> str:
    v = jpdate.normalize(value)
    if period == "as_is":
        return v
    parsed = jpdate.parse(v)
    if parsed is None:
        raise ShapingError(f"期間として解釈できない値があります: {value!r}（元号のない「26年度」などは解釈しません）")
    year, month, _ = parsed
    if period == "year":
        return f"{year:04d}"
    if month is None:
        raise ShapingError(f"年月として解釈できない値があります: {value!r}")
    return f"{year:04d}-{month:02d}"


MISSING_MARKS = {"-", "－", "‐", "―", "—", "…"}
# Units written after numbers in PDF / HTML tables. 「約」 or footnote marks are
# not stripped: such a value is reported instead of being read as exact.
_NUMBER = re.compile(r"^([△▲-]?)(\d+(?:\.\d+)?)\s*(円|件|人|か所|カ所|ヵ所|ヶ所|ケ所|箇所|%|km|世帯|回|校|園|台|点|歳)?$")


def _number(raw: str) -> Optional[float]:
    v = unicodedata.normalize("NFKC", raw).replace(",", "").strip()
    if v == "" or v in MISSING_MARKS:
        return None
    m = _NUMBER.match(v)
    if not m:
        raise ShapingError(f"数値として解釈できない値があります: {raw!r}")
    return -float(m.group(2)) if m.group(1) else float(m.group(2))


def apply_plan(rows: List[Dict[str, str]], plan: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Deterministically turn table rows into [{period, value}] sorted by period.

    probe-agent:
      role: 採用された整形計画を実データへ適用する変換処理
      capability: 根拠付きKPIダッシュボードを公開する
      system_purpose: オープンデータをKPI時系列へ変換し、公開ダッシュボードの実データを成立させる
      probe_value: AI候補から人手採用後に実際のデータ変換へ進む境界であり、MVPデータパイプラインの中核を観測できる
      element_type: core
      operation_kind: mutation
      consumers: [apps.api.app.routes.data.adopt_shaping]
      state_effects: [filesystem]
    """
    if not plan.get("feasible"):
        raise ShapingError("この整形案は実行できないとされています")
    cols = set(rows[0]) if rows else set()
    for key in ("date_column", "value_column", "filter_column"):
        col = plan.get(key) or ""
        if col and col not in cols:
            raise ShapingError(f"列が見つかりません: {col}")
    agg = plan["aggregation"]
    if agg not in AGGREGATIONS or plan["period"] not in PERIODS:
        raise ShapingError("未対応の集計方法です")
    if agg != "count_rows" and not plan.get("value_column"):
        raise ShapingError("値の列が指定されていません")
    buckets: "OrderedDict[str, float]" = OrderedDict()
    for r in rows:
        if plan.get("filter_column") and r.get(plan["filter_column"], "").strip() != plan.get("filter_value", "").strip():
            continue
        raw_date = r.get(plan["date_column"], "").strip()
        if not raw_date:
            continue
        p = _period(raw_date, plan["period"])
        if agg == "count_rows":
            buckets[p] = buckets.get(p, 0) + 1
            continue
        num = _number(r.get(plan["value_column"], ""))
        if num is None:
            continue
        buckets[p] = buckets.get(p, 0) + num if agg == "sum" else num
    points = [{"period": k, "value": v} for k, v in sorted(buckets.items())]
    if plan.get("cumulative"):
        total = 0.0
        for pt in points:
            total += pt["value"]
            pt["value"] = total
    return points


def dataset_file(datasets_dir: str, file_path: str) -> str:
    path = os.path.normpath(os.path.join(datasets_dir, file_path))
    if not path.startswith(os.path.normpath(datasets_dir) + os.sep):
        raise ShapingError("データセットのパスが不正です")
    return path
